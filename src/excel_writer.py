"""Excel output with data, validation report, and extraction log sheets."""
from __future__ import annotations
import os
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

def write_workbook(frames, validation, logs, output_path):
    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        if not frames:
            pd.DataFrame([{"Status":"No tables extracted","Next step":"Check PDF quality or install OCR dependencies."}]).to_excel(writer, index=False, sheet_name="Data")
        else:
            for i, df in enumerate(frames, start=1):
                # Keep provenance columns visible for traceability.
                df.to_excel(writer, index=False, sheet_name=f"Table_{i}"[:31])
        pd.DataFrame(validation or [{"severity":"INFO","check":"none","message":"No validation warnings were generated.","source":""}]).to_excel(writer, index=False, sheet_name="Validation")
        pd.DataFrame(logs or [{"page":"","table":"","method":"","rows":0,"columns":0}]).to_excel(writer, index=False, sheet_name="Extraction_Log")
    wb = load_workbook(output_path)
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="17365D")
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        for col in ws.columns:
            max_len = 0
            letter = get_column_letter(col[0].column)
            for cell in list(col)[:250]:
                try:
                    max_len = max(max_len, len(str(cell.value or "")))
                except Exception:
                    pass
            ws.column_dimensions[letter].width = min(max(max_len + 2, 12), 32)
        ws.row_dimensions[1].height = 28
    wb.save(output_path)
